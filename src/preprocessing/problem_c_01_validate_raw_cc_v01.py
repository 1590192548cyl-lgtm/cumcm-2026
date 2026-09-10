"""Validate the immutable Problem C source workbooks and submission templates."""

from __future__ import annotations

import json
from datetime import datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
CONFIG_PATH = REPO_ROOT / "config/problem_c_config_v01.json"


def _load_config() -> dict[str, Any]:
    return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))


def _workbook_shape(path: Path) -> dict[str, list[int]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    return {
        worksheet.title: [worksheet.max_row, worksheet.max_column]
        for worksheet in workbook.worksheets
    }


def _assert(condition: bool, message: str, checks: list[str]) -> None:
    if not condition:
        raise ValueError(message)
    checks.append(message)


def _time_key(value: Any) -> str:
    if isinstance(value, time):
        return value.strftime("%H:%M")
    text = str(value).strip()
    suffix = "+1" if text.endswith("+1") else ""
    clock = text.removesuffix("+1")
    parsed = datetime.strptime(clock, "%H:%M")
    return parsed.strftime("%H:%M") + suffix


def main() -> None:
    config = _load_config()
    paths = config["paths"]
    checks: list[str] = []
    summary: dict[str, Any] = {
        "status": "running",
        "checked_at": datetime.now().isoformat(timespec="seconds"),
        "checks": checks,
        "workbooks": {},
    }

    source_paths = {
        key: REPO_ROOT / paths[key]
        for key in ("attachment_1", "attachment_2", "attachment_3", "attachment_4")
    }
    template_paths = {
        "result1": REPO_ROOT / "data/raw/problem_c_raw_result_1_template_v01.xlsx",
        "result2": REPO_ROOT / "data/raw/problem_c_raw_result_2_template_v01.xlsx",
        "result3": REPO_ROOT / "data/raw/problem_c_raw_result_3_template_v01.xlsx",
        "result4_2": REPO_ROOT / "data/raw/problem_c_raw_result_4_2_template_v01.xlsx",
        "result4_3": REPO_ROOT / "data/raw/problem_c_raw_result_4_3_template_v01.xlsx",
    }

    for label, path in {**source_paths, **template_paths}.items():
        _assert(path.exists(), f"{label}: file exists", checks)
        summary["workbooks"][label] = {
            "path": str(path.relative_to(REPO_ROOT)),
            "size_bytes": path.stat().st_size,
            "sheets": _workbook_shape(path),
        }

    wb1 = load_workbook(source_paths["attachment_1"], read_only=True, data_only=True)
    ws1 = wb1.active
    _assert(ws1.max_row == 145 and ws1.max_column == 4, "attachment_1: 144 ten-minute rows", checks)
    _assert(
        tuple(cell.value for cell in ws1[1])
        == ("时间", "电价", "小区负载", "光伏发电预测功率"),
        "attachment_1: expected headers",
        checks,
    )
    q1_times = [row[0] for row in ws1.iter_rows(min_row=2, max_col=1, values_only=True)]
    q1_time_keys = [_time_key(value) for value in q1_times]
    _assert(len(q1_time_keys) == 144, "attachment_1: valid time cells", checks)
    _assert(len(set(q1_time_keys)) == 144, "attachment_1: no duplicate intervals", checks)

    wb2 = load_workbook(source_paths["attachment_2"], read_only=True, data_only=True)
    _assert(
        wb2.sheetnames == ["小区负载", "光伏发电实际功率"],
        "attachment_2: expected sheets",
        checks,
    )
    for sheet_name in wb2.sheetnames:
        worksheet = wb2[sheet_name]
        _assert(
            worksheet.max_row == 366 and worksheet.max_column == 145,
            f"attachment_2/{sheet_name}: 365 days by 144 intervals",
            checks,
        )
        dates = [row[0] for row in worksheet.iter_rows(min_row=2, max_col=1, values_only=True)]
        _assert(
            dates[0].date().isoformat() == "2025-01-01"
            and dates[-1].date().isoformat() == "2025-12-31",
            f"attachment_2/{sheet_name}: full 2025 coverage",
            checks,
        )

    wb3 = load_workbook(source_paths["attachment_3"], read_only=True, data_only=True)
    ws3 = wb3.active
    _assert(ws3.max_row == 1461 and ws3.max_column == 26, "attachment_3: 365 x 4 forecast issues", checks)
    issue_times = [
        row[0]
        for row in ws3.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True)
    ]
    expected_issue_times = ["0:00", "6:00", "12:00", "18:00"]
    _assert(
        all(issue_times[index : index + 4] == expected_issue_times for index in range(0, 1460, 4)),
        "attachment_3: issue times repeat at 0/6/12/18",
        checks,
    )

    wb4 = load_workbook(source_paths["attachment_4"], read_only=True, data_only=True)
    ws4 = wb4.active
    _assert(ws4.max_row == 366 and ws4.max_column == 145, "attachment_4: 365 days by 144 intervals", checks)

    expected_template_sheets = {
        "result1": ["计划购电量", "充放电量"],
        "result2": ["计划购电量", "充放电量", "紧急购电量"],
        "result3": ["计划购电量", "调整购电量", "充放电量", "紧急购电量"],
        "result4_2": ["计划购电量", "充放电量", "紧急购电量"],
        "result4_3": ["计划购电量", "调整购电量", "充放电量", "紧急购电量"],
    }
    for label, expected_sheets in expected_template_sheets.items():
        workbook = load_workbook(template_paths[label], read_only=True, data_only=False)
        _assert(workbook.sheetnames == expected_sheets, f"{label}: expected template sheets", checks)

    report_path = REPO_ROOT / paths["validation_report"]
    report_path.parent.mkdir(parents=True, exist_ok=True)
    summary["status"] = "pass"
    summary["check_count"] = len(checks)
    report_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"PASS: {len(checks)} raw-data and template checks")
    print(report_path.relative_to(REPO_ROOT))


if __name__ == "__main__":
    main()
