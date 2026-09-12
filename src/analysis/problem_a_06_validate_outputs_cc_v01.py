"""Validate Problem A CSV and XLSX outputs against the task contract."""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
CONFIG_PATH = REPO_ROOT / "config/problem_a_config_v01.json"


def _load_config() -> dict[str, Any]:
    return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))


def _require(condition: bool, message: str, checks: list[str]) -> None:
    if not condition:
        raise ValueError(message)
    checks.append(message)


def _read_csv(path: Path, expected_columns: list[str]) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(path)
    frame = pd.read_csv(path)
    if list(frame.columns) != expected_columns:
        raise ValueError(f"Unexpected columns in {path}")
    return frame


def _check_time_grid(
    frame: pd.DataFrame,
    first_time: int,
    step: int,
    expected_rows: int | None,
    label: str,
    checks: list[str],
) -> None:
    time_s = frame["time_s"].to_numpy(float)
    _require(int(time_s[0]) == first_time, f"{label}: expected first time", checks)
    _require(np.allclose(np.diff(time_s), step), f"{label}: regular time grid", checks)
    if expected_rows is not None:
        _require(len(frame) == expected_rows, f"{label}: expected row count", checks)


def _check_monotone_rows(
    values: np.ndarray,
    direction: str,
    tolerance: float,
) -> bool:
    for row in values:
        finite = row[np.isfinite(row)]
        differences = np.diff(finite)
        if direction == "increasing" and np.any(differences < -tolerance):
            return False
        if direction == "decreasing" and np.any(differences > tolerance):
            return False
    return True


def _check_workbook(
    path: Path,
    expected_sheets: list[str],
    source_frames: list[pd.DataFrame],
    checks: list[str],
) -> None:
    _require(path.exists(), f"{path.name}: workbook exists", checks)
    workbook = load_workbook(path, read_only=True, data_only=True)
    _require(workbook.sheetnames == expected_sheets, f"{path.name}: expected sheets", checks)
    for worksheet, source in zip(workbook.worksheets, source_frames, strict=True):
        rows = list(worksheet.iter_rows(min_row=1, max_row=2, values_only=True))
        _require(len(rows) == 2, f"{path.name}/{worksheet.title}: readable header and first row", checks)
        _require(
            int(rows[1][0]) == int(source.iloc[0, 0]),
            f"{path.name}/{worksheet.title}: first time matches CSV",
            checks,
        )
        last_row = next(
            worksheet.iter_rows(
                min_row=len(source) + 1,
                max_row=len(source) + 1,
                values_only=True,
            )
        )
        _require(
            int(last_row[0]) == int(source.iloc[-1, 0]),
            f"{path.name}/{worksheet.title}: last time matches CSV",
            checks,
        )
        _require(
            len(rows[0]) == len(source.columns),
            f"{path.name}/{worksheet.title}: column count matches CSV",
            checks,
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--submission-gate",
        action="store_true",
        help="Return a nonzero status when a review warning remains.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    config = _load_config()
    threshold = config["model_contract"]["drying_threshold_kg_per_kg"]
    display_threshold = config["model_contract"]["four_decimal_display_threshold_kg_per_kg"]
    results_dir = REPO_ROOT / "outputs/results"
    submissions_dir = REPO_ROOT / "outputs/submissions/problem_a"
    checks: list[str] = []
    warnings: list[str] = []

    fixed_columns = ["time_s"] + [f"r_{radius / 10:.1f}_cm" for radius in range(21)]
    shrink_columns = (
        ["time_s"] + [f"r_{radius / 10:.1f}_cm" for radius in range(20)] + ["surface"]
    )
    q1_temperature = _read_csv(
        results_dir / "problem_a_result_01_preheat_temperature_v01.csv",
        fixed_columns,
    )
    q1_moisture = _read_csv(
        results_dir / "problem_a_result_02_preheat_moisture_v01.csv",
        fixed_columns,
    )
    q2_temperature = _read_csv(
        results_dir / "problem_a_result_05_q2_temperature_v01.csv",
        fixed_columns,
    )
    q2_moisture = _read_csv(
        results_dir / "problem_a_result_06_q2_moisture_v01.csv",
        fixed_columns,
    )
    q3_moisture = _read_csv(
        results_dir / "problem_a_result_07_q3_moisture_v01.csv",
        fixed_columns,
    )
    q4_moisture = _read_csv(
        results_dir / "problem_a_result_08_q4_moisture_v01.csv",
        shrink_columns,
    )

    _check_time_grid(q1_temperature, 1, 1, 1800, "q1 temperature", checks)
    _check_time_grid(q1_moisture, 1, 1, 1800, "q1 moisture", checks)
    _check_time_grid(q2_temperature, 1, 1, 10800, "q2 temperature", checks)
    _check_time_grid(q2_moisture, 1, 1, 10800, "q2 moisture", checks)
    _check_time_grid(q3_moisture, 60, 60, None, "q3 moisture", checks)
    _check_time_grid(q4_moisture, 60, 60, None, "q4 moisture", checks)

    q1_temperature_values = q1_temperature.iloc[:, 1:].to_numpy(float)
    _require(
        np.isfinite(q1_temperature_values).all(),
        "q1 temperature: all values finite",
        checks,
    )
    _require(
        _check_monotone_rows(q1_temperature_values, "increasing", 1.0e-7),
        "q1 temperature: radial profiles increase toward surface",
        checks,
    )

    q2_temperature_values = q2_temperature.iloc[:, 1:].to_numpy(float)
    _require(
        np.isfinite(q2_temperature_values).all(),
        "q2 temperature: all values finite",
        checks,
    )
    _require(
        np.nanmin(q2_temperature_values) >= 28.0 - 1.0e-6,
        "q2 temperature: no value falls below the initial temperature",
        checks,
    )
    q2_min_radial_difference = float(
        np.nanmin(np.diff(q2_temperature_values, axis=1))
    )
    _require(
        q2_min_radial_difference >= -1.0e-2,
        (
            "q2 temperature: transient radial reversals caused by fluctuating "
            f"air temperature remain below 0.01 C (worst {q2_min_radial_difference:.6f} C)"
        ),
        checks,
    )

    for label, frame in (
        ("q1 moisture", q1_moisture),
        ("q2 moisture", q2_moisture),
        ("q3 moisture", q3_moisture),
        ("q4 moisture", q4_moisture),
    ):
        values = frame.iloc[:, 1:].to_numpy(float)
        _require(np.nanmin(values) >= 0, f"{label}: values are nonnegative", checks)
        _require(
            _check_monotone_rows(values, "decreasing", 1.0e-7),
            f"{label}: radial profiles decrease toward surface",
            checks,
        )

    q3_final_max = float(np.nanmax(q3_moisture.iloc[-1, 1:].to_numpy(float)))
    q4_final_max = float(np.nanmax(q4_moisture.iloc[-1, 1:].to_numpy(float)))
    _require(q3_final_max < threshold, "q3: final maximum is below 0.15", checks)
    _require(q4_final_max < threshold, "q4: final maximum is below 0.15", checks)
    if q3_final_max >= display_threshold:
        warnings.append(
            "Question 3 ends below 0.15 mathematically but rounds to 0.1500 at four decimals."
        )
    if q4_final_max >= display_threshold:
        warnings.append(
            "Question 4 ends below 0.15 mathematically but rounds to 0.1500 at four decimals."
        )
    warnings.append(
        "Questions 3 and 4 use different material correlations; their time difference is not a pure shrinkage effect."
    )
    warnings.append(
        "Questions 2-4 retain h and hm from Appendix 2 and hold the final air observation after 14400 s; sensitivity checks remain required."
    )

    _check_workbook(
        submissions_dir / "result1.xlsx",
        ["温度", "水分浓度"],
        [q1_temperature, q1_moisture],
        checks,
    )
    _check_workbook(
        submissions_dir / "result2.xlsx",
        ["温度", "水分浓度"],
        [q2_temperature, q2_moisture],
        checks,
    )
    _check_workbook(
        submissions_dir / "result3.xlsx",
        ["Sheet1"],
        [q3_moisture],
        checks,
    )
    _check_workbook(
        submissions_dir / "result4.xlsx",
        ["Sheet1"],
        [q4_moisture],
        checks,
    )

    report = {
        "status": "pass" if not warnings else "pass_with_warnings",
        "checked_at": datetime.now().isoformat(timespec="seconds"),
        "check_count": len(checks),
        "warning_count": len(warnings),
        "checks": checks,
        "warnings": warnings,
        "terminal_values": {
            "question_3_time_s": int(q3_moisture.iloc[-1, 0]),
            "question_3_max_moisture": q3_final_max,
            "question_4_time_s": int(q4_moisture.iloc[-1, 0]),
            "question_4_max_moisture": q4_final_max,
        },
    }
    output_path = REPO_ROOT / config["paths"]["workflow_report"]
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{report['status'].upper()}: {len(checks)} checks, {len(warnings)} warnings")
    for warning in warnings:
        print(f"WARNING: {warning}")
    print(output_path.relative_to(REPO_ROOT))
    if args.submission_gate and warnings:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
